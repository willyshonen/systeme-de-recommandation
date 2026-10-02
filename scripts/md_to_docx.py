"""
Convertit rapport_stage.md en rapport_stage.docx
Usage: python scripts/md_to_docx.py
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ─── Helpers ──────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color: str):
    """Applique une couleur de fond à une cellule."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_horizontal_rule(doc):
    """Ajoute une ligne horizontale."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pb.append(bottom)
    pPr.append(pb)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)


def apply_inline_bold_italic(run_parent, text: str):
    """Parse le texte avec **gras** et `code` inline et ajoute les runs."""
    # On parse les patterns inline : **bold**, *italic*, `code`
    pattern = re.compile(r"(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)")
    last = 0
    for m in pattern.finditer(text):
        # texte avant
        before = text[last:m.start()]
        if before:
            run = run_parent.add_run(before)
            run.font.size = Pt(11)

        if m.group(0).startswith("**"):
            run = run_parent.add_run(m.group(2))
            run.bold = True
            run.font.size = Pt(11)
        elif m.group(0).startswith("*"):
            run = run_parent.add_run(m.group(3))
            run.italic = True
            run.font.size = Pt(11)
        elif m.group(0).startswith("`"):
            run = run_parent.add_run(m.group(4))
            run.font.name = "Courier New"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)

        last = m.end()

    # texte restant
    remaining = text[last:]
    if remaining:
        run = run_parent.add_run(remaining)
        run.font.size = Pt(11)


def add_styled_paragraph(doc, text: str, style_name: str = "Normal", bold=False, italic=False, size=11, color=None, align=None, space_before=None, space_after=None):
    """Ajoute un paragraphe avec inline markdown parsé."""
    p = doc.add_paragraph()
    p.style = doc.styles[style_name] if style_name in [s.name for s in doc.styles] else doc.styles["Normal"]
    if align:
        p.alignment = align
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)

    # Strip leading › (blockquote) si présent
    text = re.sub(r"^>\s*", "", text)

    apply_inline_bold_italic(p, text)

    # Post-traitement : forcer bold/italic/color/size sur tous les runs si demandé
    for run in p.runs:
        if bold and not run.bold:
            run.bold = bold
        if italic and not run.italic:
            run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color

    return p


# ─── Conversion principale ────────────────────────────────────────────────────

def convert_md_to_docx(md_path: Path, docx_path: Path):
    doc = Document()

    # Marges
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)

    # Style de base
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    content = md_path.read_text(encoding="utf-8")
    lines = content.splitlines()

    i = 0
    in_code_block = False
    code_lines = []
    code_lang = ""

    while i < len(lines):
        line = lines[i]

        # ── Bloc de code ──────────────────────────────────────────────────────
        if line.strip().startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lang = line.strip()[3:].strip()
                code_lines = []
            else:
                # Fin du bloc — on affiche
                in_code_block = False
                if code_lang:
                    lang_p = doc.add_paragraph()
                    lang_p.paragraph_format.space_after = Pt(0)
                    lang_run = lang_p.add_run(f"  {code_lang}")
                    lang_run.font.size = Pt(9)
                    lang_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                    lang_run.italic = True

                for cl in code_lines:
                    cp = doc.add_paragraph()
                    cp.paragraph_format.left_indent = Inches(0.3)
                    cp.paragraph_format.space_before = Pt(0)
                    cp.paragraph_format.space_after = Pt(0)
                    try:
                        cp.style = doc.styles["No Spacing"]
                    except KeyError:
                        pass
                    cr = cp.add_run(cl)
                    cr.font.name = "Courier New"
                    cr.font.size = Pt(9)
                    cr.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)

                    # Fond gris clair via shading
                    pPr = cp._p.get_or_add_pPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "F5F5F5")
                    pPr.append(shd)

                doc.add_paragraph()  # espace après le bloc
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # ── Ligne vide ────────────────────────────────────────────────────────
        if not line.strip():
            i += 1
            continue

        # ── Séparateur horizontal ---  ────────────────────────────────────────
        if re.match(r"^-{3,}$", line.strip()):
            add_horizontal_rule(doc)
            i += 1
            continue

        # ── Tableau ───────────────────────────────────────────────────────────
        if line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1

            # Filtrer les séparateurs (|---|---|)
            data_rows = [r for r in table_lines if not re.match(r"^\|[\s\-\|:]+\|$", r.strip())]
            if not data_rows:
                continue

            def parse_row(r):
                cells = [c.strip() for c in r.strip().strip("|").split("|")]
                return cells

            rows = [parse_row(r) for r in data_rows]
            n_cols = max(len(r) for r in rows)

            table = doc.add_table(rows=len(rows), cols=n_cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for ri, row_data in enumerate(rows):
                for ci, cell_text in enumerate(row_data):
                    if ci >= n_cols:
                        break
                    cell = table.cell(ri, ci)
                    # Vider le contenu par défaut
                    cell.text = ""
                    p = cell.paragraphs[0]
                    apply_inline_bold_italic(p, cell_text)
                    for run in p.runs:
                        run.font.size = Pt(10)
                    if ri == 0:
                        # En-tête : fond bleu foncé, texte blanc
                        set_cell_bg(cell, "2C3E7E")
                        for run in p.runs:
                            run.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    elif ri % 2 == 0:
                        set_cell_bg(cell, "EEF0F8")

            doc.add_paragraph()
            continue

        # ── Titres # ─────────────────────────────────────────────────────────
        heading_match = re.match(r"^(#{1,6})\s+(.*)", line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            # Supprimer les émojis si trop lourds (optionnel — on les garde)
            try:
                heading = doc.add_heading(level=min(level, 6))
                heading.clear()
                run = heading.add_run(text)
                sizes = {1: 20, 2: 16, 3: 14, 4: 12, 5: 11, 6: 11}
                run.font.size = Pt(sizes.get(level, 11))
                colors = {
                    1: RGBColor(0x1A, 0x23, 0x56),
                    2: RGBColor(0x2C, 0x3E, 0x7E),
                    3: RGBColor(0x2C, 0x3E, 0x7E),
                    4: RGBColor(0x55, 0x55, 0x55),
                }
                if level in colors:
                    run.font.color.rgb = colors[level]
            except Exception:
                doc.add_paragraph(text)
            i += 1
            continue

        # ── Blockquote > ──────────────────────────────────────────────────────
        if line.strip().startswith(">"):
            text = re.sub(r"^>\s*", "", line.strip())
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "12")
            left.set(qn("w:space"), "4")
            left.set(qn("w:color"), "4472C4")
            pBdr.append(left)
            pPr.append(pBdr)
            apply_inline_bold_italic(p, text)
            for run in p.runs:
                run.font.size = Pt(10)
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # ── Liste à puces - / * ───────────────────────────────────────────────
        if re.match(r"^(\s*)[-*]\s+", line):
            indent_level = len(re.match(r"^(\s*)", line).group(1)) // 2
            text = re.sub(r"^\s*[-*]\s+", "", line)
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.25 + indent_level * 0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.clear()
            apply_inline_bold_italic(p, text)
            for run in p.runs:
                run.font.size = Pt(11)
            i += 1
            continue

        # ── Liste numérotée 1. 2. ────────────────────────────────────────────
        if re.match(r"^\d+\.\s+", line):
            text = re.sub(r"^\d+\.\s+", "", line)
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.clear()
            apply_inline_bold_italic(p, text)
            for run in p.runs:
                run.font.size = Pt(11)
            i += 1
            continue

        # ── Texte normal ─────────────────────────────────────────────────────
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        apply_inline_bold_italic(p, line.strip())
        for run in p.runs:
            run.font.size = Pt(11)
        i += 1

    doc.save(str(docx_path))
    print(f"[OK] Fichier genere : {docx_path}")


# ─── Entrée ───────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    base = Path(__file__).parent.parent
    md_path = base / "rapport_stage.md"
    docx_path = base / "rapport_stage.docx"

    if not md_path.exists():
        print(f"[ERREUR] Fichier introuvable : {md_path}")
        raise SystemExit(1)

    print(f"[...] Conversion de {md_path.name} -> {docx_path.name} ...")
    convert_md_to_docx(md_path, docx_path)
